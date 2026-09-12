from datetime import date, datetime
import io
import os
import re
import sqlite3
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import streamlit as st

# ================= 1. CẤU HÌNH GIAO DIỆN & CSS TRÀN MÀN HÌNH =================
st.set_page_config(
    page_title="EMIC QC Dashboard Pro",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
        header[data-testid="stHeader"] { display: none !important; }
        
        .main .block-container, div[data-testid="stAppViewBlockContainer"] {
            padding-top: 0.3rem !important;
            padding-bottom: 0.3rem !important;
            padding-left: 0.8rem !important;
            padding-right: 0.8rem !important;
            max-width: 100% !important;
        }

        .stApp {
            background-color: #F8FAFC;
            font-family: system-ui, -apple-system, sans-serif;
        }

        /* Styling Tabs Căn Giữa */
        .stTabs [data-baseweb="tab-list"] {
            justify-content: center !important;
            gap: 8px !important;
            background-color: transparent !important;
            padding: 2px 0px 8px 0px !important;
            border-bottom: 1px solid #E2E8F0 !important;
        }
        .stTabs [data-baseweb="tab"] {
            background-color: #94A3B8 !important;
            color: #FFFFFF !important;
            border-radius: 6px !important;
            padding: 5px 18px !important;
            border: none !important;
            font-weight: 600 !important;
            font-size: 13px !important;
        }
        .stTabs [aria-selected="true"] {
            background-color: #3B82F6 !important;
            color: #FFFFFF !important;
            box-shadow: 0 2px 4px rgba(59, 130, 246, 0.3) !important;
        }
        
        /* Card chứa biểu đồ động */
        div[data-testid="stHorizontalBlock"] > div[data-testid="stColumn"] {
            background-color: #FFFFFF;
            border-radius: 10px;
            border: 1px solid #CBD5E1;
            padding: 8px 12px !important;
            box-shadow: 0 1px 3px rgba(0,0,0,0.02);
        }

        div[data-testid="stVerticalBlock"] > div {
            gap: 0.3rem !important;
        }
        
        section[data-testid="stSidebar"] {
            background-color: #FFFFFF;
            border-right: 1px solid #E2E8F0;
        }
    </style>
""",
    unsafe_allow_html=True,
)

# Palette màu sắc phẳng chuẩn Dashboard
COLOR_SUCCESS = "#10B981"
COLOR_PRIMARY = "#3B82F6"
COLOR_DANGER = "#EF4444"
COLOR_WARNING = "#F59E0B"
COLOR_PURPLE = "#A855F7"
COLOR_TEXT = "#0F172A"

DISTINCT_COLORS = [
    "#3B82F6",
    "#10B981",
    "#F59E0B",
    "#EF4444",
    "#A855F7",
    "#06B6D4",
    "#F43F5E",
    "#84CC16",
    "#F97316",
    "#14B8A6",
]


def clean_emoji(text):
  return re.sub(r"[^\w\s\(\)\-\/\.\,\:]", "", str(text)).strip()


DB_NAME = "Report_Database.db"


@st.cache_data(ttl=30)
def load_data(tu_date, den_date):
  if not os.path.exists(DB_NAME):
    return pd.DataFrame(), pd.DataFrame()

  tu_iso = tu_date.strftime("%Y-%m-%d 00:00:00")
  den_iso = den_date.strftime("%Y-%m-%d 23:59:59")

  conn = sqlite3.connect(DB_NAME)
  df_qa32 = pd.read_sql_query(
      "SELECT * FROM tb_sap_qa32 WHERE ngay_ve_dt >= ? AND ngay_ve_dt <= ?",
      conn,
      params=(tu_iso, den_iso),
  )
  df_coois = pd.read_sql_query(
      "SELECT * FROM tb_sap_coois WHERE ngay_lenh_dt >= ? AND ngay_lenh_dt <="
      " ?",
      conn,
      params=(tu_iso, den_iso),
  )
  conn.close()
  return df_qa32, df_coois


# ================= 2. HÀM TẠO FILE EXCEL CHUẨN IN ẤN TRỔ TÀI CHUYÊN NGHIỆP =================
def generate_print_ready_excel(
    phan_he_code,
    title_clean,
    tu_date,
    den_date,
    df_monthly,
    df_plan,
    df_family,
    df_year,
):
  output = io.BytesIO()
  wb = openpyxl.Workbook()
  wb.remove(wb.active)  # Xóa sheet mặc định

  # Định dạng Styles chuẩn doanh nghiệp
  font_company = Font(name="Arial", size=10, bold=True, color="1F4E79")
  font_title = Font(name="Arial", size=15, bold=True, color="000000")
  font_subtitle = Font(name="Arial", size=9, italic=True, color="595959")
  font_section = Font(name="Arial", size=11, bold=True, color="1F4E79")
  font_header = Font(name="Arial", size=9, bold=True, color="FFFFFF")
  font_data = Font(name="Arial", size=9)

  fill_header = PatternFill(
      start_color="1F4E79", end_color="1F4E79", fill_type="solid"
  )
  fill_zebra = PatternFill(
      start_color="F9FBFD", end_color="F9FBFD", fill_type="solid"
  )

  thin_border = Border(
      left=Side(style="thin", color="D9D9D9"),
      right=Side(style="thin", color="D9D9D9"),
      top=Side(style="thin", color="D9D9D9"),
      bottom=Side(style="thin", color="D9D9D9"),
  )
  header_border = Border(
      left=Side(style="thin", color="FFFFFF"),
      right=Side(style="thin", color="FFFFFF"),
      top=Side(style="medium", color="1F4E79"),
      bottom=Side(style="medium", color="1F4E79"),
  )

  align_center = Alignment(
      horizontal="center", vertical="center", wrap_text=True
  )
  align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
  align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

  sheets_data = [
      ("TienDo_Thang", "1. TIẾN ĐỘ SẢN XUẤT THEO THÁNG", df_monthly),
      ("TongQuan_KeHoach", "2. TỔNG QUAN CHỈ TIÊU KẾ HOẠCH", df_plan),
      ("Dong_SanPham", "3. CHI TIẾT THEO DÒNG SẢN PHẨM", df_family),
      ("SanLuong_CaNam", "4. TỔNG SẢN LƯỢNG CẢ NĂM MÃ ĐẦU 5", df_year),
  ]

  for sheet_name, section_title, df_table in sheets_data:
    ws = wb.create_sheet(title=sheet_name)

    # Cấu hình in ấn khổ giấy A4 Nằm Ngang (Landscape Fit Page)
    ws.views.sheetView[0].showGridLines = True
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Header Thông Tin Công Ty & Báo Cáo
    ws["A1"] = "TỔNG CÔNG TY THIẾT BỊ ĐIỆN EMIC - PHÒNG QUẢN LÝ CHẤT LƯỢNG (QC)"
    ws["A1"].font = font_company

    ws["A2"] = f"BÁO CÁO TỔNG HỢP SẢN XUẤT & QUẢN LÝ CHẤT LƯỢNG - {title_clean.upper()}"
    ws["A2"].font = font_title

    ws["A3"] = (
        f"Giai đoạn: Từ ngày {tu_date.strftime('%d/%m/%Y')} đến ngày"
        f" {den_date.strftime('%d/%m/%Y')} | Thời điểm xuất dữ liệu:"
        f" {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )
    ws["A3"].font = font_subtitle

    # Tiêu đề bảng
    ws["A5"] = section_title
    ws["A5"].font = font_section

    start_row = 7

    # Ghi Header Bảng
    for col_idx, col_name in enumerate(df_table.columns, start=1):
      cell = ws.cell(row=start_row, column=col_idx, value=col_name)
      cell.font = font_header
      cell.fill = fill_header
      cell.alignment = align_center
      cell.border = header_border

    # Ghi Dữ Liệu & Formatting Số/Phần Trăm
    for row_idx, row_data in enumerate(df_table.values, start=start_row + 1):
      is_even = row_idx % 2 == 0
      row_fill = fill_zebra if is_even else PatternFill(fill_type=None)

      for col_idx, val in enumerate(row_data, start=1):
        cell = ws.cell(row=row_idx, column=col_idx)
        col_name_lower = str(df_table.columns[col_idx - 1]).lower()

        if isinstance(val, (int, float, np.number)):
          if "%" in col_name_lower or "tỷ lệ" in col_name_lower:
            cell.value = float(val) / 100.0 if val > 1.0 else float(val)
            cell.number_format = "0.0%"
          else:
            cell.value = float(val)
            cell.number_format = (
                "#,##0" if float(val).is_integer() else "#,##0.00"
            )
          cell.alignment = align_right
        else:
          cell.value = str(val)
          cell.alignment = (
              align_center
              if col_idx == 1 or len(str(val)) < 10
              else align_left
          )

        cell.font = font_data
        if row_fill.fill_type:
          cell.fill = row_fill
        cell.border = thin_border

    # Căn chỉnh tự động độ rộng cột không bị cắt chữ / lỗi ###
    for col in ws.columns:
      max_len = 0
      col_letter = get_column_letter(col[0].column)
      for cell in col:
        if cell.row >= start_row and cell.value is not None:
          max_len = max(max_len, len(str(cell.value)))
      ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

  wb.save(output)
  return output.getvalue()


# ================= 3. SIDEBAR BỘ LỌC DỮ LIỆU =================
st.sidebar.title("🎛️ BỘ LỌC DỮ LIỆU")
col_tu, col_den = st.sidebar.columns(2)
with col_tu:
  tu_date = st.date_input("Từ ngày", date(datetime.now().year, 1, 1))
with col_den:
  den_date = st.date_input("Đến ngày", date.today())

st.sidebar.markdown("---")
if st.sidebar.button("🔄 Cập Nhật Dữ Liệu", use_container_width=True):
  st.cache_data.clear()
  st.rerun()

df_qa32, df_coois = load_data(tu_date, den_date)

# ================= 4. NAVIGATION TABS CĂN GIỮA =================
tab_vat_tu, tab_co_khi, tab_tuti, tab_cong_to = st.tabs([
    "📋 Báo Cáo Vật Tư",
    "⚙️ Báo Cáo Cơ Khí",
    "🔌 Báo Cáo TU/TI",
    "⚡ Báo Cáo Công Tơ",
])

# ================= 5. TAB 1: BÁO CÁO VẬT TƯ (BIỂU ĐỒ ĐỘNG PLOTLY) =================
with tab_vat_tu:
  if df_qa32.empty:
    st.info("💡 Chưa có dữ liệu QA32 trong khoảng thời gian đã chọn.")
  else:
    months_labels = [f"T{i}" for i in range(1, 13)]
    ud01_m, ud02_m, ud03_m, uninspected_m = [0] * 12, [0] * 12, [0] * 12, [0] * 12
    ft_qty_m, by_inspected_m, by_uninspected_m = (
        [0.0] * 12,
        [0.0] * 12,
        [0.0] * 12,
    )
    total_ca_block, total_ft_all = 0.0, 0.0
    top_block_dict = {}

    for _, r in df_qa32.iterrows():
      try:
        m_idx = (
            datetime.strptime(str(r["ngay_ve_dt"]).split()[0], "%Y-%m-%d").month
            - 1
        )
      except:
        m_idx = 0
      if not (0 <= m_idx < 12):
        m_idx = 0

      st_clean = (
          str(r["xac_nhan_sap"]).strip().upper().replace(" ", "")
          if pd.notna(r["xac_nhan_sap"])
          else ""
      )
      ft_val = (
          float(r["ft_qty"])
          if ("ft_qty" in r and pd.notna(r["ft_qty"]))
          else 0.0
      )
      by_val = (
          float(r["by_sample"])
          if ("by_sample" in r and pd.notna(r["by_sample"]))
          else 0.0
      )
      ca_val = (
          float(r["ca_qty"])
          if ("ca_qty" in r and pd.notna(r["ca_qty"]))
          else 0.0
      )

      ft_qty_m[m_idx] += ft_val
      total_ft_all += ft_val
      total_ca_block += ca_val

      is_uninspected = (
          "CHƯA" in st_clean
          or not st_clean
          or st_clean in ["NAN", "NONE", "❌CHƯAXN"]
      )
      is_ud02 = any(
          k in st_clean for k in ["02", "UD2", "ĐẶCNHƯỢNG", "DACNHUONG"]
      )
      is_ud03 = any(
          k in st_clean
          for k in [
              "03",
              "UD3",
              "TRẢLẠI",
              "TRALAI",
              "TỪCHỐI",
              "TUCHOI",
              "KHÔNG",
              "KHONG",
          ]
      )
      is_ud01 = any(k in st_clean for k in ["01", "UD1", "ĐẠT", "DAT"]) and not (
          is_ud02 or is_ud03
      )

      if is_uninspected:
        uninspected_m[m_idx] += 1
        by_uninspected_m[m_idx] += by_val
      else:
        by_inspected_m[m_idx] += by_val
        if is_ud02:
          ud02_m[m_idx] += 1
        elif is_ud03:
          ud03_m[m_idx] += 1
        else:
          ud01_m[m_idx] += 1

      ma_vt_str = str(r["ma_vt"]).strip() if pd.notna(r["ma_vt"]) else ""
      ten_vt_str = str(r["ten_vt"]).strip() if pd.notna(r["ten_vt"]) else ""
      ncc_str = str(r["ncc"]).strip() if pd.notna(r["ncc"]) else ""

      if (
          is_ud02
          or is_ud03
          or ca_val > 0
          or (st_clean and not is_ud01 and not is_uninspected)
      ):
        key = (ma_vt_str, ncc_str)
        if key not in top_block_dict:
          top_block_dict[key] = {
              "ma_vt": ma_vt_str,
              "ten_vt": ten_vt_str,
              "ncc": ncc_str,
              "ud02": 0,
              "ud03": 0,
              "ca_block": 0.0,
              "ft_total": 0.0,
          }
        if is_ud02:
          top_block_dict[key]["ud02"] += 1
        if is_ud03:
          top_block_dict[key]["ud03"] += 1
        top_block_dict[key]["ca_block"] += ca_val
        top_block_dict[key]["ft_total"] += ft_val

    col1, col2 = st.columns([2.2, 1.0])

    with col1:
      fig1 = make_subplots(specs=[[{"secondary_y": True}]])
      fig1.add_trace(
          go.Bar(
              x=months_labels,
              y=ud01_m,
              name="UD 01 (Đạt)",
              marker_color=COLOR_SUCCESS,
          ),
          secondary_y=False,
      )
      fig1.add_trace(
          go.Bar(
              x=months_labels,
              y=ud02_m,
              name="UD 02 (Đặc nhượng)",
              marker_color=COLOR_WARNING,
          ),
          secondary_y=False,
      )
      fig1.add_trace(
          go.Bar(
              x=months_labels,
              y=ud03_m,
              name="UD 03 (Trả lại)",
              marker_color=COLOR_DANGER,
          ),
          secondary_y=False,
      )

      total_by = [by_inspected_m[i] + by_uninspected_m[i] for i in range(12)]
      fig1.add_trace(
          go.Scatter(
              x=months_labels,
              y=total_by,
              name="Số mẫu phải kiểm (BY)",
              line=dict(color=COLOR_PURPLE, width=2, dash="dash"),
          ),
          secondary_y=True,
      )
      fig1.add_trace(
          go.Scatter(
              x=months_labels,
              y=ft_qty_m,
              name="Tổng số hàng về (FT)",
              line=dict(color=COLOR_PRIMARY, width=2),
          ),
          secondary_y=True,
      )

      fig1.update_layout(
          title=dict(
              text="BÁO CÁO SỐ LƯỢNG LỆNH KIỂM & TỔNG VẬT TƯ VỀ / SỐ MẪU KIỂM",
              font=dict(size=12, color=COLOR_TEXT),
              x=0.5,
              xanchor="center",
          ),
          barmode="stack",
          margin=dict(l=10, r=10, t=35, b=10),
          height=240,
          paper_bgcolor="rgba(0,0,0,0)",
          plot_bgcolor="rgba(0,0,0,0)",
          legend=dict(
              orientation="h",
              yanchor="top",
              y=-0.20,
              xanchor="center",
              x=0.5,
              font=dict(size=9),
          ),
      )
      fig1.update_yaxes(
          title_text="← Số Lượng Lệnh",
          secondary_y=False,
          showgrid=True,
          gridcolor="#E2E8F0",
      )
      fig1.update_yaxes(
          title_text="Vật Tư / Mẫu [Log] →",
          type="log",
          secondary_y=True,
          showgrid=False,
      )

      st.plotly_chart(
          fig1, use_container_width=True, config={"displayModeBar": False}
      )

    with col2:
      ok_cnt = max(0.0, total_ft_all - total_ca_block)
      fig2 = go.Figure(
          data=[
              go.Pie(
                  labels=["Vật tư Đạt", "Bị Block (Lỗi)"],
                  values=[ok_cnt, total_ca_block],
                  hole=0.5,
                  marker_colors=[COLOR_SUCCESS, COLOR_DANGER],
                  textinfo="label+percent",
              )
          ]
      )
      fig2.update_layout(
          title=dict(
              text="TỶ LỆ VẬT TƯ ĐẠT VS BỊ BLOCK LỖI",
              font=dict(size=12, color=COLOR_TEXT),
              x=0.5,
              xanchor="center",
          ),
          margin=dict(l=10, r=10, t=35, b=10),
          height=240,
          paper_bgcolor="rgba(0,0,0,0)",
          showlegend=False,
      )
      st.plotly_chart(
          fig2, use_container_width=True, config={"displayModeBar": False}
      )

    sorted_blocks = sorted(
        top_block_dict.values(),
        key=lambda x: (x["ud03"] + x["ud02"], x["ca_block"], x["ft_total"]),
        reverse=True,
    )
    if sorted_blocks:
      df_block = pd.DataFrame(sorted_blocks)
      df_block["Tổng SL Block (CA) / SL Về"] = df_block.apply(
          lambda r: f"{r['ca_block']:,.0f} / {r['ft_total']:,.0f}", axis=1
      )
      df_block = df_block[[
          "ma_vt",
          "ten_vt",
          "ncc",
          "ud02",
          "ud03",
          "Tổng SL Block (CA) / SL Về",
      ]]
      df_block.columns = [
          "Mã Vật Tư",
          "Tên Vật Tư",
          "Nhà Cung Cấp",
          "Số Lượt UD 02",
          "Số Lượt UD 03",
          "Tổng SL Block (CA) / SL Về",
      ]
      st.dataframe(df_block, use_container_width=True, hide_index=True)


# ================= 6. HÀM CHUNG CHO CÁC TAB COOIS (BIỂU ĐỒ ĐỘNG PLOTLY 2x2) =================
def render_coois_tab_layout(phan_he_code, title_text):
  df_sub = (
      df_coois[df_coois["phan_he"] == phan_he_code]
      if not df_coois.empty
      else pd.DataFrame()
  )

  if df_sub.empty:
    st.info(f"💡 Chưa có dữ liệu sản xuất cho phân hệ {title_text}.")
    return

  months_labels = [f"T{i}" for i in range(1, 13)]
  m_comp_qty, m_uncomp_qty = [0.0] * 12, [0.0] * 12
  m_tot_orders, m_uncomp_orders = [0] * 12, [0] * 12
  tot_qty_all, deliv_qty_all = 0.0, 0.0

  for _, r in df_sub.iterrows():
    try:
      m_idx = (
          datetime.strptime(str(r["ngay_lenh_dt"]).split()[0], "%Y-%m-%d").month
          - 1
      )
    except:
      m_idx = 0
    if not (0 <= m_idx < 12):
      m_idx = 0

    sl_t, sl_h = float(r["sl_tong"]), float(r["sl_ht"])
    uncomp_q = max(0.0, sl_t - sl_h)

    tot_qty_all += sl_t
    deliv_qty_all += sl_h
    m_comp_qty[m_idx] += sl_h
    m_uncomp_qty[m_idx] += uncomp_q

    m_tot_orders[m_idx] += 1
    if sl_h < sl_t:
      m_uncomp_orders[m_idx] += 1

  title_clean = clean_emoji(title_text)

  # HÀNG 1: BIỂU ĐỒ TIẾN ĐỘ THÁNG & DONUT TỔNG QUAN
  col1, col2 = st.columns([2.2, 1.0])

  with col1:
    fig1 = make_subplots(specs=[[{"secondary_y": True}]])
    m_comp_orders = [m_tot_orders[i] - m_uncomp_orders[i] for i in range(12)]

    fig1.add_trace(
        go.Bar(
            x=months_labels,
            y=m_comp_orders,
            name="Lệnh Hoàn Thành",
            marker_color=COLOR_PRIMARY,
        ),
        secondary_y=False,
    )
    fig1.add_trace(
        go.Bar(
            x=months_labels,
            y=m_uncomp_orders,
            name="Lệnh Chưa Xong",
            marker_color=COLOR_DANGER,
        ),
        secondary_y=False,
    )
    fig1.add_trace(
        go.Bar(
            x=months_labels,
            y=m_comp_qty,
            name="SL Hoàn Thành",
            marker_color=COLOR_SUCCESS,
        ),
        secondary_y=True,
    )
    fig1.add_trace(
        go.Bar(
            x=months_labels,
            y=m_uncomp_qty,
            name="SL Chưa Xong",
            marker_color=COLOR_WARNING,
        ),
        secondary_y=True,
    )

    fig1.update_layout(
        title=dict(
            text=f"TIẾN ĐỘ SẢN XUẤT - {title_clean}",
            font=dict(size=12, color=COLOR_TEXT),
            x=0.5,
            xanchor="center",
        ),
        barmode="stack",
        margin=dict(l=10, r=10, t=35, b=10),
        height=240,
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        legend=dict(
            orientation="h",
            yanchor="top",
            y=-0.20,
            xanchor="center",
            x=0.5,
            font=dict(size=9),
        ),
    )
    fig1.update_yaxes(
        title_text="← Tổng Lệnh",
        secondary_y=False,
        showgrid=True,
        gridcolor="#E2E8F0",
    )
    fig1.update_yaxes(
        title_text="Số Lượng Giao [Log] →",
        type="log",
        secondary_y=True,
        showgrid=False,
    )

    st.plotly_chart(
        fig1, use_container_width=True, config={"displayModeBar": False}
    )

  with col2:
    rem_qty_all = max(0.0, tot_qty_all - deliv_qty_all)
    pct_deliv = (deliv_qty_all / tot_qty_all * 100) if tot_qty_all > 0 else 0
    pct_rem = 100.0 - pct_deliv if tot_qty_all > 0 else 0.0

    fig2 = go.Figure(
        data=[
            go.Pie(
                labels=["Hoàn thành", "Chưa xong"],
                values=[deliv_qty_all, rem_qty_all],
                hole=0.5,
                marker_colors=[COLOR_SUCCESS, COLOR_WARNING],
                textinfo="label+percent",
            )
        ]
    )
    fig2.update_layout(
        title=dict(
            text="TỶ LỆ HOÀN THÀNH TỔNG QUAN",
            font=dict(size=12, color=COLOR_TEXT),
            x=0.5,
            xanchor="center",
        ),
        margin=dict(l=10, r=10, t=35, b=10),
        height=240,
        paper_bgcolor="rgba(0,0,0,0)",
        showlegend=False,
        annotations=[
            dict(
                text=f"TỔNG KẾ HOẠCH<br><b>{int(tot_qty_all):,}</b>",
                x=0.5,
                y=0.5,
                font_size=10,
                showarrow=False,
            )
        ],
    )
    st.plotly_chart(
        fig2, use_container_width=True, config={"displayModeBar": False}
    )

  # HÀNG 2: BỘ LỌC DÒNG SP & 2 BIỂU ĐỒ BÊN DƯỚI
  sub_5 = (
      df_sub[
          df_sub["ma_tp"]
          .astype(str)
          .str.split(".")
          .str[0]
          .str.lstrip("0")
          .str.startswith("5")
      ].copy()
      if not df_sub.empty
      else pd.DataFrame()
  )

  raw_fams = (
      [
          str(x).strip()
          for x in sub_5["mat_prefix"].unique()
          if pd.notna(x)
          and str(x).strip()
          and str(x).strip().lower() not in ["none", "nan"]
      ]
      if not sub_5.empty
      else []
  )
  clean_fams = sorted(list(set(raw_fams)))
  available_fams = ["Tất cả dòng sản phẩm"] + clean_fams

  sel_fam = st.selectbox(
      "🎯 Chọn Dòng SP (Đầu 5):", available_fams, key=f"cb_{phan_he_code}"
  )

  col3, col4 = st.columns([1, 1])

  # --- DƯỚI TRÁI: SẢN LƯỢNG DÒNG SP ---
  with col3:
    m3_qty = [0.0] * 12
    if not sub_5.empty:
      sub_5_df = sub_5.copy()
      sub_5_df["month"] = pd.to_datetime(
          sub_5_df["ngay_lenh_dt"], errors="coerce"
      ).dt.month
      sub_filtered = (
          sub_5_df[sub_5_df["mat_prefix"] == sel_fam]
          if sel_fam != "Tất cả dòng sản phẩm"
          else sub_5_df
      )
      for _, r in sub_filtered.iterrows():
        m_val = r["month"]
        if pd.notna(m_val) and 1 <= int(m_val) <= 12:
          m3_qty[int(m_val) - 1] += float(r["sl_ht"])

    fig3 = go.Figure()
    fig3.add_trace(
        go.Bar(
            x=months_labels,
            y=m3_qty,
            name="SL Sản Xuất",
            marker_color=COLOR_PRIMARY,
            text=[f"{int(v):,}" if v > 0 else "" for v in m3_qty],
            textposition="auto",
        )
    )
    fig3.update_layout(
        title=dict(
            text=f"SẢN LƯỢNG - DÒNG: {clean_emoji(sel_fam)}",
            font=dict(size=12, color=COLOR_TEXT),
            x=0.5,
            xanchor="center",
        ),
        margin=dict(l=10, r=10, t=35, b=10),
        height=240,
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        showlegend=False,
    )
    fig3.update_yaxes(
        title_text="SL Hoàn Thành [Log]",
        type="log",
        showgrid=True,
        gridcolor="#E2E8F0",
    )

    st.plotly_chart(
        fig3, use_container_width=True, config={"displayModeBar": False}
    )

  # --- DƯỚI PHẢI: TỔNG SẢN LƯỢNG MÃ ĐẦU 5 ---
  with col4:
    if not sub_5.empty:
      summary_fams = (
          sub_5.groupby("mat_prefix")[["sl_ht"]].sum().reset_index()
      )
      fams_x = [
          str(val)
          for val in summary_fams["mat_prefix"].tolist()
          if pd.notna(val)
          and str(val).strip()
          and str(val).strip().lower() not in ["none", "nan"]
      ]
      if not fams_x:
        fams_x, deliv_fams = ["Trống"], [0.0]
      else:
        summary_fams = summary_fams[summary_fams["mat_prefix"].isin(fams_x)]
        fams_x = summary_fams["mat_prefix"].tolist()
        deliv_fams = summary_fams["sl_ht"].values
    else:
      fams_x, deliv_fams = ["Không có SP"], [0.0]

    bar_colors = [
        DISTINCT_COLORS[i % len(DISTINCT_COLORS)] for i in range(len(fams_x))
    ]

    fig4 = go.Figure()
    fig4.add_trace(
        go.Bar(
            x=fams_x,
            y=deliv_fams,
            marker_color=bar_colors,
            text=[f"{int(v):,}" if v > 0 else "" for v in deliv_fams],
            textposition="auto",
        )
    )
    fig4.update_layout(
        title=dict(
            text="TỔNG SẢN LƯỢNG CẢ NĂM CÁC MÃ ĐẦU 5",
            font=dict(size=12, color=COLOR_TEXT),
            x=0.5,
            xanchor="center",
        ),
        margin=dict(l=10, r=10, t=35, b=10),
        height=240,
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        showlegend=False,
    )
    fig4.update_yaxes(
        title_text="Số Lượng SP [Log]",
        type="log",
        showgrid=True,
        gridcolor="#E2E8F0",
    )

    st.plotly_chart(
        fig4, use_container_width=True, config={"displayModeBar": False}
    )

  # ================= 7. BẢNG SỐ LIỆU TỔNG HỢP & NÚT XUẤT EXCEL IN ẤN =================
  st.markdown("---")

  pct_orders_m = [
      ((m_comp_orders[i] / m_tot_orders[i]) * 100.0)
      if m_tot_orders[i] > 0
      else 0.0
      for i in range(12)
  ]
  pct_qty_m = [
      ((m_comp_qty[i] / (m_comp_qty[i] + m_uncomp_qty[i])) * 100.0)
      if (m_comp_qty[i] + m_uncomp_qty[i]) > 0
      else 0.0
      for i in range(12)
  ]

  total_m3 = sum(m3_qty)
  pct_fam_contrib = [
      ((m3_qty[i] / total_m3) * 100.0) if total_m3 > 0 else 0.0
      for i in range(12)
  ]

  total_yr = sum(deliv_fams)
  pct_yr_share = [
      ((v / total_yr) * 100.0) if total_yr > 0 else 0.0 for v in deliv_fams
  ]

  df_monthly_summary = pd.DataFrame({
      "Tháng": months_labels,
      "Lệnh Hoàn Thành": m_comp_orders,
      "Lệnh Chưa Xong": m_uncomp_orders,
      "Tỷ Lệ HT Lệnh (%)": pct_orders_m,
      "SL Hoàn Thành": [int(v) for v in m_comp_qty],
      "SL Chưa Xong": [int(v) for v in m_uncomp_qty],
      "Tỷ Lệ HT SL (%)": pct_qty_m,
  })

  df_plan_summary = pd.DataFrame({
      "Chỉ Tiêu": ["Tổng Kế Hoạch", "Đã Giao Hoàn Thành", "Còn Lại Chưa Xong"],
      "Số Lượng": [int(tot_qty_all), int(deliv_qty_all), int(rem_qty_all)],
      "Tỷ Lệ Cơ Cấu (%)": [
          100.0,
          (deliv_qty_all / tot_qty_all * 100.0) if tot_qty_all > 0 else 0.0,
          (rem_qty_all / tot_qty_all * 100.0) if tot_qty_all > 0 else 0.0,
      ],
  })

  df_family_summary = pd.DataFrame({
      "Tháng": months_labels,
      f"SL Sản Xuất ({sel_fam})": [int(v) for v in m3_qty],
      "Tỷ Lệ Đóng Góp Tháng (%)": pct_fam_contrib,
  })

  df_year_summary = pd.DataFrame({
      "Mã / Dòng SP": fams_x,
      "Tổng SL Hoàn Thành Cả Năm": [int(v) for v in deliv_fams],
      "Tỷ Lệ Cơ Cấu (%)": pct_yr_share,
  })

  col_hdr, col_btn = st.columns([2.5, 1.2])
  with col_hdr:
    st.markdown(
        f"### 📑 BẢNG SỐ LIỆU TỔNG HỢP & TỶ LỆ HIỆU CHỈNH -"
        f" {title_clean.upper()}"
    )

  excel_bytes = generate_print_ready_excel(
      phan_he_code,
      title_clean,
      tu_date,
      den_date,
      df_monthly_summary,
      df_plan_summary,
      df_family_summary,
      df_year_summary,
  )

  with col_btn:
    st.download_button(
        label="📥 XUẤT BÁO CÁO EXCEL (CHUẨN IN A4)",
        data=excel_bytes,
        file_name=f"BaoCao_EMIC_{phan_he_code}_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary",
    )

  # HIỂN THỊ 4 BẢNG SỐ LIỆU MÀU SẮC DẠNG PROGRESS BAR
  t_col1, t_col2 = st.columns([1.6, 1.0])
  with t_col1:
    st.markdown("##### 1. Tiến Độ Sản Xuất Theo Tháng")
    st.dataframe(
        df_monthly_summary,
        column_config={
            "Tỷ Lệ HT Lệnh (%)": st.column_config.ProgressColumn(
                "Tỷ Lệ HT Lệnh", format="%.1f%%", min_value=0, max_value=100
            ),
            "Tỷ Lệ HT SL (%)": st.column_config.ProgressColumn(
                "Tỷ Lệ HT SL", format="%.1f%%", min_value=0, max_value=100
            ),
            "SL Hoàn Thành": st.column_config.NumberColumn(
                "SL Hoàn Thành", format="%d"
            ),
            "SL Chưa Xong": st.column_config.NumberColumn(
                "SL Chưa Xong", format="%d"
            ),
        },
        use_container_width=True,
        hide_index=True,
    )

  with t_col2:
    st.markdown("##### 2. Tổng Quan Chỉ Tiêu Kế Hoạch")
    st.dataframe(
        df_plan_summary,
        column_config={
            "Tỷ Lệ Cơ Cấu (%)": st.column_config.ProgressColumn(
                "Tỷ Lệ Cơ Cấu", format="%.1f%%", min_value=0, max_value=100
            ),
            "Số Lượng": st.column_config.NumberColumn("Số Lượng", format="%d"),
        },
        use_container_width=True,
        hide_index=True,
    )

  t_col3, t_col4 = st.columns([1.0, 1.0])
  with t_col3:
    st.markdown(f"##### 3. Chi Tiết Sản Lượng Dòng SP ({sel_fam})")
    st.dataframe(
        df_family_summary,
        column_config={
            "Tỷ Lệ Đóng Góp Tháng (%)": st.column_config.ProgressColumn(
                "Tỷ Lệ Đóng Góp Tháng",
                format="%.1f%%",
                min_value=0,
                max_value=100,
            ),
            f"SL Sản Xuất ({sel_fam})": st.column_config.NumberColumn(
                f"SL Sản Xuất ({sel_fam})", format="%d"
            ),
        },
        use_container_width=True,
        hide_index=True,
    )

  with t_col4:
    st.markdown("##### 4. Tổng Sản Lượng Cả Năm Các Mã Đầu 5")
    st.dataframe(
        df_year_summary,
        column_config={
            "Tỷ Lệ Cơ Cấu (%)": st.column_config.ProgressColumn(
                "Tỷ Lệ Cơ Cấu", format="%.1f%%", min_value=0, max_value=100
            ),
            "Tổng SL Hoàn Thành Cả Năm": st.column_config.NumberColumn(
                "Tổng SL Cả Năm", format="%d"
            ),
        },
        use_container_width=True,
        hide_index=True,
    )


# ================= 8. RENDER NỘI DUNG CÁC TAB =================
with tab_co_khi:
  render_coois_tab_layout("CO_KHI", "⚙️ BÁO CÁO CƠ KHÍ (LỆNH 3012)")
with tab_tuti:
  render_coois_tab_layout("TU_TI", "🔌 BÁO CÁO TUTI (LỆNH 3011)")
with tab_cong_to:
  render_coois_tab_layout("CONG_TO", "⚡ BÁO CÁO CÔNG TƠ (LỆNH 3013, 3016)")